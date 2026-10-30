"""Tests for demo/symfony8/bin/generate-stress-docx.py (REQ-TEST-004).

The generator writes into a temporary directory: the curated demo fixture
(public/demo/stress-styles.docx) is never touched.
"""

from __future__ import annotations

import importlib.util
import runpy
from pathlib import Path

import pytest
from docx import Document
from docx.enum.section import WD_ORIENT
from PIL import Image

SCRIPT = Path(__file__).resolve().parents[1] / "demo" / "symfony8" / "bin" / "generate-stress-docx.py"


def _load_module():
    spec = importlib.util.spec_from_file_location("generate_stress_docx", SCRIPT)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture()
def generator(tmp_path, monkeypatch):
    module = _load_module()
    monkeypatch.setattr(module, "OUT", tmp_path / "out" / "stress-styles.docx")
    monkeypatch.setattr(module, "ASSETS", tmp_path / "assets")
    return module


def test_default_paths_point_at_the_demo_public_dir():
    module = _load_module()
    assert module.OUT.parts[-3:] == ("public", "demo", "stress-styles.docx")
    assert module.ASSETS.parts[-3:] == ("public", "demo", "_stress_assets")


def test_make_images_writes_three_pngs(generator):
    paths = generator._make_images()

    assert set(paths) == {"logo", "chart", "photo"}
    sizes = {name: Image.open(path).size for name, path in paths.items()}
    assert sizes == {"logo": (640, 160), "chart": (720, 360), "photo": (400, 280)}


def test_make_images_falls_back_to_default_font(generator, monkeypatch):
    original = generator.ImageFont.truetype
    requested = []

    def missing_dejavu(font, *args, **kwargs):
        # Pillow's load_default() itself calls truetype() with an in-memory font: keep that working.
        if isinstance(font, str):
            requested.append(font)
            raise OSError("font not installed")
        return original(font, *args, **kwargs)

    monkeypatch.setattr(generator.ImageFont, "truetype", missing_dejavu)

    paths = generator._make_images()

    assert requested == ["/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"]
    assert all(path.is_file() for path in paths.values())


def test_make_images_uses_dejavu_when_installed(generator, monkeypatch):
    original = generator.ImageFont.truetype
    load_default = generator.ImageFont.load_default
    requested = []

    def dejavu(font, size, *args, **kwargs):
        if not isinstance(font, str):
            # load_default() re-enters truetype() with Pillow's embedded font.
            return original(font, size, *args, **kwargs)
        requested.append((font, size))
        return load_default(size=size)

    monkeypatch.setattr(generator.ImageFont, "truetype", dejavu)

    generator._make_images()

    assert requested == [
        ("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 36),
        ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 18),
    ]


def test_set_section_columns_reuses_existing_cols_element(generator):
    doc = Document()
    section = doc.sections[0]
    for existing in section._sectPr.findall(generator.qn("w:cols")):
        section._sectPr.remove(existing)

    generator._set_section_columns(section, num=2)
    assert len(section._sectPr.findall(generator.qn("w:cols"))) == 1
    generator._set_section_columns(section, num=3, space_twips=200)

    cols = section._sectPr.findall(generator.qn("w:cols"))
    assert len(cols) == 1
    assert cols[0].get(generator.qn("w:num")) == "3"
    assert cols[0].get(generator.qn("w:space")) == "200"


def test_build_writes_a_valid_stress_document(generator):
    out = generator.build()

    assert out == generator.OUT
    assert out.is_file() and out.stat().st_size > 0

    doc = Document(str(out))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    for expected in ("1. Document map", "4. Tables under pressure", "5. Embedded images", "9. Closing checklist"):
        assert expected in headings

    assert len(doc.inline_shapes) >= 3
    assert len(doc.tables) >= 2
    assert any(section.orientation == WD_ORIENT.LANDSCAPE for section in doc.sections)
    assert any(section.orientation == WD_ORIENT.PORTRAIT for section in doc.sections)
    assert "Callout" in [style.name for style in doc.styles]


def test_set_run_font_handles_serif_and_mono_east_asian_fonts(generator):
    doc = Document()
    paragraph = doc.add_paragraph()

    serif = paragraph.add_run("serif")
    generator._set_run_font(serif, "Liberation Serif")
    mono = paragraph.add_run("mono")
    generator._set_run_font(mono, "Liberation Mono", 9, bold=False, italic=True)

    qn = generator.qn
    assert serif._element.rPr.rFonts.get(qn("w:eastAsia")) == "Noto Serif CJK JP"
    assert mono._element.rPr.rFonts.get(qn("w:eastAsia")) == "Liberation Mono"
    assert mono.font.italic is True


def test_running_as_script_writes_into_the_patched_location(tmp_path, monkeypatch, capsys):
    # Run the real __main__ block against a copy of the script so OUT/ASSETS resolve under tmp_path.
    bin_dir = tmp_path / "demo" / "bin"
    bin_dir.mkdir(parents=True)
    copy = bin_dir / "generate-stress-docx.py"
    copy.write_text(SCRIPT.read_text(encoding="utf-8"), encoding="utf-8")

    runpy.run_path(str(copy), run_name="__main__")

    out = tmp_path / "demo" / "public" / "demo" / "stress-styles.docx"
    assert out.is_file()
    assert f"Wrote {out}" in capsys.readouterr().out
